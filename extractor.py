import streamlit as st
from docx import Document
from fpdf import FPDF
import tempfile
import re
from datetime import datetime,timedelta
import json
import requests

def get_req(key):
    return st.secrets[key]

GITHUB_TOKEN=get_req("GITHUB_TOKEN")
REPO_NAME=get_req("REPO")

if not REPO_NAME or not GITHUB_TOKEN:
    raise Exception("Environment not configured properly")

HEADERS={
    "Authorization": f"token {GITHUB_TOKEN}",
    "Accept": f"application/vnd.github.v3.raw"
}

def load_json_from_github(path,repo=REPO_NAME):
    url=f"https://api.github.com/repos/{repo}/contents/{path}"
    r=requests.get(url,headers=HEADERS)
    if r.status_code!=200:
        raise Exception(f"Github API failed: {r.status_code} - {r.text}")
    return json.loads(r.text)

TIME_SLOTS = [
    "9.00-9.55", "10.00-10.55", "11.00-11.55", "12.00-12.55",
    "1.00-1.55", "2.00-2.55", "3.00-3.55", "4.00-4.55"
]

DAYS= ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]

def load_batch_groups(path='batch_groups.json'):
    return load_json_from_github(path)
    
BATCH_GROUPS = load_batch_groups()

def load_dates(path="dates.json"):
    data=load_json_from_github(path)
    semester_start = datetime.strptime(data["semester_start"], "%Y-%m-%d").date()
    semester_end = datetime.strptime(data["semester_end"], "%Y-%m-%d").date()
    #single days
    holidays=set()
    for date_str in data.get("single_days",[]):
        holidays.add(datetime.strptime(date_str,"%Y-%m-%d").date())
    #ranges
    for rng in data.get("ranges",[]):
        start=datetime.strptime(rng["start"],"%Y-%m-%d").date()
        end=datetime.strptime(rng["end"],"%Y-%m-%d").date()
        while(start<=end):
            holidays.add(start)
            start+=timedelta(days=1)
    return semester_start, semester_end, holidays

SEMESTER_START, SEMESTER_END, HOLIDAYS = load_dates()
#print(HOLIDAYS)

def get_group_for_batch(batch_code):
    for group, members in BATCH_GROUPS.items():
        if batch_code in members:
            return group
    return None

def split_batches(batch_str):
    #if comma-separated, use that
    if ',' in batch_str:
        return [b.strip() for b in batch_str.split(',')]
    #if contains space, split by space
    elif ' ' in batch_str:
        return [b.strip() for b in batch_str.split(' ')]
    #else, split by capital letter + number pattern (e.g., B7B8 -> [B7,B8])
    return re.findall(r'[A-Z]\d+', batch_str)

def extract_timetable(docx_file, batch_code):
    doc = Document(docx_file)
    timetable = {day: [] for day in DAYS}
    group_code = get_group_for_batch(batch_code)

    for table in doc.tables:
        rows = table.rows
        i = 0
        while i < len(rows) - 1:
            row1 = rows[i].cells
            row2 = rows[i + 1].cells
            day = row1[0].text.strip()[:3]
            if day not in timetable:
                i += 1
                continue
            for col in range(1, min(len(TIME_SLOTS) + 1, len(row1))):
                slot = TIME_SLOTS[col - 1]
                entry1 = row1[col].text.strip()
                entry2 = row2[col].text.strip()
                added = False
                for line in entry1.split('\n') + entry2.split('\n'):
                    if not line:
                        continue
                    batch_part = line.split('-')[0].strip()
                    batches = split_batches(batch_part)

                    if batch_code in batches or (group_code and batch_part == group_code):
                        timetable[day].append((slot, line))
                        added = True
                        break
                if not added:
                    timetable[day].append((slot, None))
            i += 2
    return timetable

def generate_pdf(timetable, batch_code):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.set_title(f"{batch_code} Timetable")
    pdf.cell(200, 10, txt=f"Timetable for Batch: {batch_code}", ln=True, align="C")
    pdf.ln(10)
    for day in timetable:
        pdf.set_font("Arial", "B", 12)
        pdf.cell(200, 10, txt=day, ln=True)
        pdf.set_font("Arial", size=12)
        for slot, entry in timetable[day]:  
            if(slot=="1.00-1.55"):
                line = f"{slot}: Lunch"
            else:
                line = f"{slot}: {entry if entry else 'No class'}"
            pdf.cell(200, 8, txt=line, ln=True)
        pdf.ln(5)
    tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf.output(tmp_file.name)
    return tmp_file.name

DAY_TO_INDEX = {"Mon": 0, "Tue": 1, "Wed": 2, "Thu": 3, "Fri": 4, "Sat": 5}

START_DATE = datetime.combine(SEMESTER_START, datetime.min.time())
until_str= SEMESTER_END.strftime("%Y%m%dT235959")

def generate_ics(timetable):
    tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".ics")
    with open(tmp_file.name, "w") as icsfile:
        icsfile.write("BEGIN:VCALENDAR\nVERSION:2.0\nPRODID:-//My Timetable//EN\n")
        for day, slots in timetable.items():
            day_index = DAY_TO_INDEX[day]
            event_date = START_DATE + timedelta(days=day_index)
            for time_range, class_info in slots:
                if class_info is None:
                    continue
                start_str, end_str = [t.replace('.', ':') for t in time_range.split('-')]
                start_time = datetime.strptime(start_str, "%H:%M")
                end_time = datetime.strptime(end_str, "%H:%M")

                # adjust for PM slots
                if 1 <= start_time.hour <= 4:
                    start_time = start_time.replace(hour=start_time.hour + 12)
                    end_time = end_time.replace(hour=end_time.hour + 12)

                dtstart = event_date.replace(hour=start_time.hour, minute=start_time.minute)
                dtend = event_date.replace(hour=end_time.hour, minute=end_time.minute)
                dtstart_str = dtstart.strftime("%Y%m%dT%H%M%S")
                dtend_str = dtend.strftime("%Y%m%dT%H%M%S")

                icsfile.write("BEGIN:VEVENT\n")
                icsfile.write(f"SUMMARY:{class_info}\n")
                icsfile.write(f"DTSTART:{dtstart_str}\n")
                icsfile.write(f"DTEND:{dtend_str}\n")
                icsfile.write(f"RRULE:FREQ=WEEKLY;UNTIL={until_str}\n")

                for holiday in HOLIDAYS:
                    if holiday.weekday() == dtstart.weekday():
                        ex_dt=datetime.combine(holiday, dtstart.time())
                        ex_str=ex_dt.strftime("%Y%m%dT%H%M%S")
                        # print(f"Adding EXDATE for {class_info} on {ex_dt} because {holiday} matches weekday {dtstart.strftime('%A')}")
                        icsfile.write(f"EXDATE:{ex_str}\n")
                icsfile.write("END:VEVENT\n")
        icsfile.write("END:VCALENDAR\n")
    return tmp_file.name